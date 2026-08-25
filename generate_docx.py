from docx import Document
from pathlib import Path

files = [
    'index.html',
    'events.html',
    'about.html',
    'contact.html',
    'studentzone.html',
    'style.css',
    'script.js'
]

doc = Document()
doc.add_heading('Society Website Source Code', level=1)
for file_name in files:
    path = Path(file_name)
    if not path.exists():
        continue
    doc.add_heading(file_name, level=2)
    doc.add_paragraph(f'File path: {str(path.resolve())}')
    doc.add_paragraph('')
    content = path.read_text(encoding='utf-8')
    paragraph = doc.add_paragraph()
    for line in content.splitlines():
        paragraph.add_run(line)
        paragraph.add_run('\n')
    doc.add_page_break()

output = Path('Society Website Source.docx')
doc.save(output)
print(output.name)
